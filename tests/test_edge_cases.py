#!/usr/bin/env python3
"""Edge-case torture tests for the built-in engine.

Builds a set of awkward documents (merged cells, nested tables, long tables that
span pages, multi-section files, oversized images, tabs, deep list levels, empty
documents ...) and verifies that every one converts without crashing and that the
resulting PDFs look sane.

    python3 tests/test_edge_cases.py [-v]
"""

from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, ROOT)

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from make_samples import add_numbering, add_field, apply_numbering, make_png, shade_cell  # noqa: E402

VERBOSE = "-v" in sys.argv or "--verbose" in sys.argv
_results = []


def check(name: str, condition: bool, detail: str = "") -> bool:
    _results.append((name, bool(condition), detail))
    if VERBOSE or not condition:
        print(f"  [{'PASS' if condition else 'FAIL'}] {name}"
              + (f" -- {detail}" if detail and not condition else ""))
    return bool(condition)


def set_cell_borders(cell, **edges) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    xml = "".join(
        f'<w:{edge} w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        for edge, (size, color) in edges.items())
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}>{xml}</w:tcBorders>'))


# --------------------------------------------------------------------------- #
# documents
# --------------------------------------------------------------------------- #


def doc_empty(path: str) -> None:
    Document().save(path)


def doc_only_table(path: str) -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(1, 1).text = "D"
    doc.save(path)


def doc_merged_and_nested(path: str) -> None:
    doc = Document()
    doc.add_paragraph("合并单元格与嵌套表格 / merged + nested")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "跨两列 / span 2"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "跨两行 / span rows"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(0, 2).text = "R0C2"
    table.cell(1, 1).text = "R1C1"
    table.cell(2, 1).text = "R2C1"
    table.cell(1, 2).text = "R1C2"
    table.cell(2, 2).text = "R2C2"
    # nested table in the last cell
    nested = table.cell(2, 2).add_table(rows=2, cols=2)
    nested.style = "Table Grid"
    nested.cell(0, 0).text = "嵌套"
    nested.cell(1, 1).text = "inner"
    doc.save(path)


def doc_long_table(path: str) -> None:
    doc = Document()
    doc.add_paragraph("长表格分页 / long table pagination")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(("序号", "名称 Name", "说明 Description")):
        cell = table.cell(0, index)
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
        shade_cell(cell, "D9E2F3")
    # mark the first row as a repeating header
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>'))
    for index in range(1, 61):
        row = table.add_row()
        row.cells[0].text = str(index)
        row.cells[1].text = f"项目 item {index}"
        row.cells[2].text = "描述文字 " * 4
    doc.save(path)


def doc_sections_landscape(path: str) -> None:
    doc = Document()
    doc.add_paragraph("第一节：纵向 / section 1 portrait")
    doc.add_paragraph("正文内容。" * 30)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    doc.add_paragraph("第二节：横向 / section 2 landscape")
    doc.add_paragraph("横向页面内容。" * 40)
    doc.save(path)


def doc_oversized_image(path: str) -> None:
    doc = Document()
    doc.add_paragraph("超大图片应被缩放到页面内 / oversized image must be scaled down")
    doc.add_picture(io.BytesIO(make_png(2400, 1600)), width=Cm(40))
    doc.add_paragraph("图片之后 / after the image")
    doc.save(path)


def doc_odd_formatting(path: str) -> None:
    doc = Document()
    doc.add_heading("格式边界 / formatting edge cases", level=1)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = Pt(30)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    paragraph.paragraph_format.left_indent = Cm(1)
    paragraph.paragraph_format.right_indent = Cm(1.5)
    paragraph.add_run("悬挂缩进 + 固定行距 + 右缩进。" * 3)
    tabbed = doc.add_paragraph("列1\t列2\t列3\nafter newline")
    tabbed.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    long_word = doc.add_paragraph("https://example.com/" + "a" * 220 + " end")
    long_word.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    small = doc.add_paragraph()
    run = small.add_run("上标 superscript")
    run.font.superscript = True
    small.add_run(" 与 ")
    run = small.add_run("下标 subscript")
    run.font.subscript = True
    small.add_run(" 以及 ")
    run = small.add_run("删除线 strike")
    run.font.strike = True
    # hyperlink
    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="rIdLink">'
        f'<w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr>'
        f'<w:t>外部链接 external link</w:t></w:r></w:hyperlink>',
        )
    paragraph._p.append(hyperlink)
    para_break = doc.add_paragraph("这一页在段落属性中分页。")
    para_break.paragraph_format.page_break_before = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("分页之后 / after page breaks")
    doc.save(path)


def doc_deep_lists(path: str) -> None:
    doc = Document()
    add_numbering(doc, num_id=70, abstract_id=80, fmt="decimal",
                  texts=["%1.", "%2)", "%3."])
    add_numbering(doc, num_id=71, abstract_id=81, fmt="upperRoman",
                  texts=["%1.", "%2."])
    for level, text in ((0, "第一层 level 1"), (1, "第二层 level 2"), (2, "第三层 level 3"),
                        (0, "回到第一层 back to 1"), (1, "第二层 again")):
        paragraph = doc.add_paragraph()
        paragraph.add_run(text)
        apply_numbering(paragraph, 70, level)
    for text in ("Roman 一", "Roman 二"):
        paragraph = doc.add_paragraph()
        paragraph.add_run(text)
        apply_numbering(paragraph, 71, 0)
    # list inside a table cell
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "单元格内列表 / list in cell"
    for text in ("cell bullet 1", "cell bullet 2"):
        cell_paragraph = table.cell(0, 1).add_paragraph()
        cell_paragraph.text = text
        apply_numbering(cell_paragraph, 70, 0)
    doc.save(path)


def doc_pct_table_borderless(path: str) -> None:
    doc = Document()
    doc.add_paragraph("百分比宽度无边框表格 / 100% width, no borders")
    table = doc.add_table(rows=3, cols=3)
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(
        f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>'))
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"r{r}c{c}"
    set_cell_borders(table.cell(1, 1), top=(12, "FF0000"), bottom=(6, "0000FF"),
                     left=(12, "FF0000"), right=(6, "0000FF"))
    doc.save(path)


def doc_headers_tables(path: str) -> None:
    doc = Document()
    doc.add_paragraph("页眉页脚包含表格与页码 / header & footer with table and fields")
    section = doc.sections[0]
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(16))
    table.cell(0, 0).text = "左页眉"
    table.cell(0, 1).text = "右页眉"
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("第 ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" 页")
    doc.save(path)


CASES = [
    ("empty", doc_empty, {}),
    ("only_table", doc_only_table, {}),
    ("merged_nested", doc_merged_and_nested, {"期望": ["跨两列", "R2C2", "嵌套", "inner"]}),
    ("long_table", doc_long_table, {"期望": ["项目 item 60"], "最少页数": 2}),
    ("sections", doc_sections_landscape, {"期望": ["横向页面内容"], "最少页数": 2}),
    ("oversized_image", doc_oversized_image, {"期望": ["图片之后"], "图片": True}),
    ("odd_formatting", doc_odd_formatting, {"期望": ["删除线", "外部链接", "分页之后"], "最少页数": 2}),
    ("deep_lists", doc_deep_lists, {"期望": ["第三层", "Roman 二"], "编号": ["1.", "1)", "I.", "II."]}),
    ("pct_table", doc_pct_table_borderless, {"期望": ["r2c2"]}),
    ("headers_tables", doc_headers_tables, {"期望": ["左页眉", "右页眉", "第 1 页"]}),
]


def main() -> int:
    from d2p.docxmodel import read_docx
    from d2p.renderer import PdfRenderer

    work = tempfile.mkdtemp(prefix="d2p-edge-")
    try:
        renderer = PdfRenderer()
        for name, builder, expectations in CASES:
            source = os.path.join(work, f"{name}.docx")
            target = os.path.join(work, f"{name}.pdf")
            try:
                builder(source)
            except Exception as exc:
                check(f"{name}: 生成失败 / build failed", False, repr(exc))
                continue
            try:
                model = read_docx(source)
                stats = renderer.render(model, target)
            except Exception as exc:
                import traceback

                check(f"{name}: 转换失败 / conversion failed", False,
                      f"{exc}\n{traceback.format_exc()}")
                continue
            check(f"{name}: 生成 PDF / produced PDF",
                  os.path.exists(target) and os.path.getsize(target) > 800)
            check(f"{name}: 至少 1 页 / at least 1 page", stats.pages >= 1)
            if "最少页数" in expectations:
                check(f"{name}: 至少 {expectations['最少页数']} 页 / expected pages",
                      stats.pages >= expectations["最少页数"], f"got {stats.pages}")

            import pymupdf

            with pymupdf.open(target) as pdf:
                text = " ".join(" ".join(page.get_text().split()) for page in pdf)
                if expectations.get("图片"):
                    check(f"{name}: 图片已嵌入 / image embedded",
                          any(len(page.get_images()) for page in pdf))
                check(f"{name}: 每一页都有内容或为空白 / pages opened fine", pdf.page_count >= 1)
            for needle in expectations.get("期望", []):
                check(f"{name}: 包含 {needle!r}", needle in text,
                      text[:200] if needle not in text else "")
            for marker in expectations.get("编号", []):
                check(f"{name}: 编号 {marker!r}", marker in text, text[:200])
            for warning in model.warnings:
                if VERBOSE:
                    print(f"      ! {name}: {warning}")
    finally:
        shutil.rmtree(work, ignore_errors=True)

    failed = [n for n, ok, _ in _results if not ok]
    print("-" * 52)
    print(f"通过 / passed: {len(_results) - len(failed)}/{len(_results)}")
    for name in failed:
        print(f"  x {name}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
